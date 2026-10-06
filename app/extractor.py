"""PPTX → `curso.json`: lo que entra al motor.

Produce el mismo formato que `videos/csm-curso/datos/curso.json`: por lámina, sus formas
con nombre (texto por párrafo), las notas del orador, las frases de la narración y las
imágenes. Es la tarea 0.1 de PLATAFORMA.md: el extractor que no estaba en ninguna rama.
"""

import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# Velocidad medida en el curso csm: 4 571 palabras locutadas en 1 909 s por Carlos
# (eleven_v3), 41 láminas (`csm-curso/datos/tiempos-voz.json`).
PALABRAS_POR_SEGUNDO = 2.394
# Lo que csm.py añade a cada lámina: el título entra 1,0 s antes y 1,3 s de respiro al final.
RELLENO_POR_LAMINA = 2.3

# Afirmaciones que en contenido normativo necesitan su fuente a la vista (PLATAFORMA §3.4).
_NORMATIVO = re.compile(
    r"\b(?:Ley|Decreto|Resoluci[oó]n|Circular|Art[ií]culo|Art\.)\s+(?:No\.?\s*)?\d[\d\.]*(?:\s+de\s+\d{4})?"
    r"|\d+(?:[\.,]\d+)?\s?%"
    r"|\d[\d\.]*\s+SMMLV"
    r"|\d[\d\.,]*\s+(?:metros?|kil[oó]metros(?:\s+por\s+hora)?|km/h|segundos|minutos|horas|años|veh[ií]culos|unidades)\b"
    r"|\$\s?\d[\d\.,]*",
    re.IGNORECASE,
)


def frases(texto: str) -> list[str]:
    """Parte la narración igual que csm.py: tras punto, interrogación, exclamación o dos puntos."""
    return [x for x in re.split(r"(?<=[\.\?\!:])\s+", (texto or "").strip()) if x.strip()]


def segundos(texto: str) -> float:
    palabras = len((texto or "").split())
    return round(palabras / PALABRAS_POR_SEGUNDO + (RELLENO_POR_LAMINA if palabras else 0), 1)


def afirmaciones_normativas(texto: str) -> list[str]:
    vistas = []
    for m in _NORMATIVO.finditer(texto or ""):
        cita = m.group(0).strip()
        if cita not in vistas:
            vistas.append(cita)
    return vistas


def _formas(shapes):
    """Recorre las formas, también dentro de grupos."""
    for s in shapes:
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _formas(s.shapes)
        else:
            yield s


def _nombre_imagen(lamina, forma) -> str:
    """El nombre del archivo en ppt/media/ (p. ej. `image1.jpg`), como en el curso.json de csm."""
    try:
        return Path(str(lamina.part.related_part(forma._element.blip_rId).partname)).name
    except Exception:
        return f"imagen.{forma.image.ext}"


def leer_pptx(ruta: Path) -> list[dict]:
    pres = Presentation(str(ruta))
    laminas = []
    for n, lamina in enumerate(pres.slides, start=1):
        formas, imagenes, tablas, graficos = {}, [], [], []
        for s in _formas(lamina.shapes):
            if getattr(s, "has_table", False) and s.has_table:
                filas = [[c.text.strip() for c in fila.cells] for fila in s.table.rows]
                tablas.append([f for f in filas if any(f)])
            if getattr(s, "has_chart", False) and s.has_chart:
                graficos.append(_grafico(s.chart))
            if s.has_text_frame:
                parrafos = [p.text.strip() for p in s.text_frame.paragraphs if p.text.strip()]
                if parrafos:
                    nombre, i = s.name, 2
                    while nombre in formas:  # dos formas con el mismo nombre: no se pisa ninguna
                        nombre, i = f"{s.name} ({i})", i + 1
                    formas[nombre] = parrafos
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                imagenes.append((s.width * s.height, _nombre_imagen(lamina, s)))

        notas = ""
        if lamina.has_notes_slide:
            notas = lamina.notes_slide.notes_text_frame.text.strip()

        imagenes.sort(reverse=True)
        laminas.append({
            "n": n,
            "formas": formas,
            "notas": notas,
            "frases": frases(notas),
            "foto": imagenes[0][1] if imagenes else None,
            "icono": imagenes[-1][1] if len(imagenes) > 1 else None,
            "tablas": tablas,
            "graficos": graficos,
        })
    return laminas


def _grafico(chart) -> dict:
    """Lo que dice un gráfico: título, tipo, categorías y series con sus valores."""
    titulo = None
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            titulo = chart.chart_title.text_frame.text.strip() or None
    except Exception:  # un gráfico raro no impide leer el resto de la lámina
        pass
    try:
        categorias = [str(c) for c in chart.plots[0].categories] if len(chart.plots) else []
    except Exception:
        categorias = []
    series = []
    for plot in chart.plots:
        for serie in plot.series:
            try:
                valores = [v for v in serie.values]
            except Exception:
                valores = []
            series.append({"nombre": serie.name, "valores": valores})
    return {"titulo": titulo, "tipo": str(chart.chart_type).split(" ")[0].lower(), "categorias": categorias,
            "series": series}


def guardar_imagenes(ruta: Path, carpeta: Path) -> int:
    """Copia las imágenes del PPTX a `carpeta`, con el mismo nombre que `foto` e `icono` en curso.json."""
    carpeta.mkdir(parents=True, exist_ok=True)
    guardadas = 0
    for lamina in Presentation(str(ruta)).slides:
        for s in _formas(lamina.shapes):
            if s.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            destino = carpeta / _nombre_imagen(lamina, s)
            if not destino.exists():
                destino.write_bytes(s.image.blob)
                guardadas += 1
    return guardadas


# Nombres de forma que marcan el título y la sección de una lámina. Incluyen los que pone
# PowerPoint por defecto («Título 1», «Title 1») y los del curso de moto («title», «section»).
_TITULO = ("title", "titulo", "título", "cover-title")
_SECCION = ("section", "seccion", "sección", "modulo", "módulo")


def _forma(lamina: dict, prefijos: tuple[str, ...]) -> str | None:
    for nombre, parrafos in lamina.get("formas", {}).items():
        if parrafos and nombre.lower().startswith(prefijos):
            return parrafos[0].strip()
    return None


def titulo_lamina(lamina: dict) -> str:
    """El título de la lámina: su forma «title» si la tiene; si no, su primer texto."""
    titulo = _forma(lamina, _TITULO)
    if titulo:
        return titulo[:90]
    for parrafos in lamina.get("formas", {}).values():
        if parrafos:
            return parrafos[0][:90]
    return f"Lámina {lamina['n']}"


def seccion_lamina(lamina: dict) -> str | None:
    """«APERTURA», «MÓDULO 1»…: la sección que marca la lámina en una forma con nombre."""
    return _forma(lamina, _SECCION)


def agrupar(laminas: list[dict], objetivo: float = 120.0) -> list[dict]:
    """Propuesta de videos. Si el PPTX marca secciones, un video por sección; si no, por duración."""
    return _por_seccion(laminas) or _por_duracion(laminas, objetivo)


def _por_seccion(laminas: list[dict]) -> list[dict] | None:
    """Agrupa las láminas seguidas que comparten sección. Las que no la marcan (una portada,
    por ejemplo) se unen a la sección siguiente, o a la anterior si están al final."""
    secciones = [seccion_lamina(l) for l in laminas]
    marcadas = [s for s in secciones if s]
    if len(marcadas) < len(laminas) / 2 or len(set(marcadas)) < 2:
        return None

    grupos: list[tuple[str, list[dict]]] = []
    sueltas: list[dict] = []
    for l, s in zip(laminas, secciones):
        if s is None:
            sueltas.append(l)
        elif grupos and grupos[-1][0] == s:
            grupos[-1][1].extend(sueltas + [l])
            sueltas = []
        else:
            grupos.append((s, sueltas + [l]))
            sueltas = []
    if sueltas:
        grupos[-1][1].extend(sueltas)

    salida = []
    for i, (s, ls) in enumerate(grupos, start=1):
        # El título de la sección es el de su primera lámina que sí la marca.
        primera = next(l for l in ls if seccion_lamina(l) == s)
        nombre = s[:1].upper() + s[1:].lower()
        salida.append({"clave": f"v{i:02d}", "titulo": f"{nombre} · {titulo_lamina(primera)}", "laminas": [l["n"] for l in ls]})
    return salida


def _por_duracion(laminas: list[dict], objetivo: float) -> list[dict]:
    """Láminas seguidas hasta rondar `objetivo` segundos de narración."""
    grupos, actual, dur = [], [], 0.0
    for l in laminas:
        s = segundos(l["notas"])
        if actual and dur + s > objetivo * 1.25:
            grupos.append(actual)
            actual, dur = [], 0.0
        actual.append(l["n"])
        dur += s
    if actual:
        grupos.append(actual)
    por_n = {l["n"]: l for l in laminas}
    return [
        {"clave": f"v{i:02d}", "titulo": titulo_lamina(por_n[g[0]]), "laminas": g}
        for i, g in enumerate(grupos, start=1)
    ]
