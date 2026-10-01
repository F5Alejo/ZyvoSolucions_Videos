import io
import json
import shutil
import subprocess
from pathlib import Path

import numpy as np
import pytest
import soundfile as sf
from pptx import Presentation
from pptx.util import Inches

hay_ffmpeg = shutil.which("ffmpeg") is not None


def pptx_con_foto() -> bytes:
    """Dos láminas con notas; la segunda trae una foto."""
    from PIL import Image
    foto = io.BytesIO()
    Image.new("RGB", (640, 480), (30, 80, 150)).save(foto, "PNG")
    pres = Presentation()
    for i, (titulo, notas) in enumerate([
        ("Riesgo vial", "El riesgo vial se gestiona. Lo exige la Ley 1503 de 2011."),
        ("Conducción defensiva", "Anticipa lo que harán los demás."),
    ]):
        s = pres.slides.add_slide(pres.slide_layouts[5])
        s.shapes.title.text = titulo
        s.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = "Una viñeta"
        if i == 1:
            foto.seek(0)
            s.shapes.add_picture(foto, Inches(5), Inches(2), Inches(3))
        s.notes_slide.notes_text_frame.text = notas
    buf = io.BytesIO()
    pres.save(buf)
    return buf.getvalue()


class VozDePrueba:
    """Un tono con la duración que tendría la frase a la velocidad medida en csm."""
    extension = ".wav"

    def __init__(self, voz):
        pass

    def generar(self, texto, destino):
        from app.extractor import PALABRAS_POR_SEGUNDO
        seg = max(0.3, len(texto.split()) / PALABRAS_POR_SEGUNDO)
        t = np.arange(int(seg * 24000)) / 24000
        sf.write(str(destino), (0.3 * np.sin(2 * np.pi * 220 * t)).astype(np.float32), 24000)


# ── Normalizar ───────────────────────────────────────────────────────────────

def test_normalizar_lee_normas_cifras_y_siglas():
    from motor.normalizar import para_voz
    assert para_voz("Lo exige la Ley 1503 de 2011: el 30 % de los casos.") == \
        "Lo exige la Ley mil quinientos tres de dos mil once: el treinta por ciento de los casos."
    assert para_voz("Multa de $ 1.000.000 o 15 SMMLV.") == \
        "Multa de un millón de pesos o quince salarios mínimos mensuales legales vigentes."
    assert para_voz("A 80 km/h, 2,5 segundos.") == "A ochenta kilómetros por hora, dos coma cinco segundos."
    assert para_voz("El SG-SST y el SST.") == \
        "El sistema de gestión de seguridad y salud en el trabajo y el ese ese te."
    assert para_voz("Decreto No. 1079 de 2015. ¿Aplica? No.").startswith("Decreto número mil setenta y nueve")
    assert para_voz("¿Aplica? No.") == "¿Aplica? No."
    assert para_voz("el 1.º de enero") == "el primero de enero"
    # La marca puede cambiar cómo se dice una sigla.
    assert para_voz("El PESV.", {"PESV": "plan estratégico de seguridad vial"}) == "El plan estratégico de seguridad vial."


# ── Subtítulos ───────────────────────────────────────────────────────────────

def test_subtitulos_parten_frases_largas_y_conservan_el_tiempo(tmp_path):
    from motor import subtitulos
    larga = "Conducir a la defensiva es anticipar lo que harán los demás conductores, peatones y ciclistas en la vía."
    cues = subtitulos.cues([(1.0, 3.0, "Hola."), (3.5, 9.5, larga)])
    assert cues[0] == (1.0, 3.0, "Hola.")
    partes = cues[1:]
    assert len(partes) >= 2 and all(len(l) <= 42 for _, _, t in partes for l in t.split("\n"))
    assert partes[0][0] == 3.5 and abs(partes[-1][1] - 9.5) < 1e-9
    assert " ".join(t.replace("\n", " ") for _, _, t in partes) == larga

    subtitulos.escribir_vtt(cues, tmp_path / "a.vtt")
    subtitulos.escribir_srt(cues, tmp_path / "a.srt")
    vtt = (tmp_path / "a.vtt").read_text(encoding="utf-8")
    assert vtt.startswith("WEBVTT") and "00:00:01.000 --> 00:00:03.000\nHola." in vtt
    assert (tmp_path / "a.srt").read_text(encoding="utf-8").startswith("1\n00:00:01,000 --> 00:00:03,000\nHola.")


# ── Escenas ──────────────────────────────────────────────────────────────────

def test_estilo_usa_el_bloque_video_o_deduce_uno_legible(datos_copia):
    from app import datos
    from motor import escenas
    m = datos.marcas()
    e = escenas.estilo(m["riskmann"])
    assert e["fondo"] == "#020202" and e["acento"] == "#C8951A"
    for id_ in ("fegir", "sofu", "yezid-ricaurte"):  # sin bloque «video»: se deduce de la paleta
        e = escenas.estilo(m[id_])
        assert escenas.contraste(e["fondo"], e["texto"]) >= 4.5
        assert escenas.contraste(e["fondo"], e["acento"]) >= 3 or e["acento"] == e["texto"]


def test_vista_elige_el_tipo_de_escena(tmp_path):
    from motor import escenas
    (tmp_path / "image1.png").write_bytes(b"png")
    l = {"n": 3, "formas": {"Título": ["Conducción defensiva"], "Texto": ["Anticipa", "Distancia"]}, "foto": "image1.png"}
    assert escenas.vista(l, 0, 3, "Video", tmp_path, "Curso")["tipo"] == "portada"
    v = escenas.vista(l, 1, 3, "Video", tmp_path)
    assert v["tipo"] == "imagen" and v["titulo"] == "Conducción defensiva" and v["vinetas"] == ["Anticipa", "Distancia"]
    assert escenas.vista({**l, "foto": "no-existe.png"}, 1, 3, "Video", tmp_path)["tipo"] == "lista"
    html = escenas.html(v, escenas.estilo({"paleta": [{"hex": "#000000"}, {"hex": "#FFCC00"}]}), borrador=True)
    assert "Conducción defensiva" in html and "BORRADOR" in html


# ── Imágenes del PPTX ────────────────────────────────────────────────────────

def test_al_crear_el_trabajo_se_guardan_sus_imagenes(datos_copia):
    from app import taller
    t = taller.desde_pptx("x.pptx", pptx_con_foto(), "Con foto")
    foto = t["laminas"][1]["foto"]
    assert foto and (datos_copia / "trabajos" / t["id"] / "media" / foto).exists()


# ── Audio ────────────────────────────────────────────────────────────────────

@pytest.mark.skipif(not hay_ffmpeg, reason="hace falta ffmpeg")
def test_audio_queda_a_menos_14_lufs(tmp_path):
    from motor import audio
    tono = tmp_path / "tono.wav"
    t = np.arange(48000 * 3) / 48000
    sf.write(str(tono), (0.02 * np.sin(2 * np.pi * 440 * t)).astype(np.float32), 48000)  # muy bajo
    pista = tmp_path / "pista.wav"
    audio.armar_pista([(1.0, tono), (5.0, tono)], 9.0, pista)
    assert abs(sf.info(str(pista)).duration - 9.0) < 0.01
    logrado = audio.normalizar(pista, tmp_path / "norma.wav")
    m = audio.medir(tmp_path / "norma.wav")
    assert abs(logrado - audio.LUFS) <= 0.3 and float(m["input_tp"]) <= audio.PICO


# ── De punta a punta ─────────────────────────────────────────────────────────

@pytest.mark.skipif(not hay_ffmpeg, reason="hace falta ffmpeg")
def test_de_un_pptx_sale_un_mp4_listo_para_publicar(datos_copia, monkeypatch):
    pytest.importorskip("playwright")
    from app import taller
    from motor import produccion, voz
    monkeypatch.setitem(voz.PROVEEDORES, "Kokoro", VozDePrueba)

    t = taller.desde_pptx("x.pptx", pptx_con_foto(), "Curso de prueba")
    taller.ajustar(t, "riskmann", "kokoro-dora", ["16:9"])
    avisos = []
    informe = produccion.producir(t, t["videos"][0]["clave"], avisar=lambda p, x: avisos.append(x))

    salida = datos_copia / "trabajos" / t["id"] / "salida" / t["videos"][0]["clave"]
    mp4 = salida / f"{t['videos'][0]['clave']}.mp4"
    assert mp4.exists() and not (salida / "tmp").exists()
    fallas = [c for c in informe["chequeos"] if c["ok"] is False]
    assert not fallas, fallas
    assert avisos == sorted(avisos) and avisos[-1] < 1

    info = json.loads(subprocess.run(["ffprobe", "-v", "error", "-show_streams", "-of", "json", str(mp4)],
                                     capture_output=True, text=True).stdout)
    v = next(s for s in info["streams"] if s["codec_type"] == "video")
    assert (v["width"], v["height"], v["codec_name"]) == (1920, 1080, "h264")

    frases = sum(len(l["frases"]) for l in t["laminas"])
    vtt = (salida / f"{t['videos'][0]['clave']}.vtt").read_text(encoding="utf-8")
    assert vtt.count("-->") >= frases and "Ley 1503 de 2011" in vtt  # el subtítulo no se normaliza
    assert json.loads((salida / "qa.json").read_text(encoding="utf-8"))["voz"] == "kokoro-dora"

    # En GitHub Actions se guarda el resultado como artefacto (variable GUARDAR_MUESTRA): el video,
    # sus subtítulos y su informe de calidad. La voz es un tono de prueba, no una voz real.
    import os
    destino = os.environ.get("GUARDAR_MUESTRA")
    if destino:
        Path(destino).mkdir(parents=True, exist_ok=True)
        clave = t["videos"][0]["clave"]
        for nombre in (f"{clave}.mp4", f"{clave}.vtt", f"{clave}.srt", "qa.json"):
            shutil.copy(salida / nombre, Path(destino) / f"muestra-{nombre}")


def test_voz_sin_clave_avisa_antes_de_empezar(datos_copia, monkeypatch):
    from app import taller
    from motor import produccion, voz
    monkeypatch.delenv("ELEVENLABS_API_KEY", raising=False)
    t = taller.desde_pptx("x.pptx", pptx_con_foto(), "Sin clave")
    taller.ajustar(t, "riskmann", "jerome", ["16:9"])  # sin voice_id
    with pytest.raises(voz.VozNoDisponible):
        produccion.producir(t, t["videos"][0]["clave"])


# ── Desde la API ─────────────────────────────────────────────────────────────

def _crear(c) -> str:
    r = c.post("/api/trabajos", files={"archivo": ("x.pptx", pptx_con_foto())}, data={"nombre": "Curso de prueba"})
    assert r.status_code == 201
    return r.json()["id"]


@pytest.mark.skipif(not hay_ffmpeg, reason="hace falta ffmpeg")
def test_producir_desde_la_api(datos_copia, monkeypatch):
    pytest.importorskip("playwright")
    from fastapi.testclient import TestClient

    from app.main import app
    from motor import cola, voz
    monkeypatch.setitem(voz.PROVEEDORES, "Kokoro", VozDePrueba)
    monkeypatch.setattr(voz, "disponible", lambda v: None)
    c = TestClient(app)

    id_ = _crear(c)
    c.patch(f"/api/trabajos/{id_}", json={"marca": "riskmann", "voz": "kokoro-dora", "formatos": ["16:9"]})
    p = c.get(f"/api/trabajos/{id_}/produccion").json()
    assert p["voz_falta"] is None and p["videos"] == {"v01": None}

    r = c.post(f"/api/trabajos/{id_}/producir/v01")
    assert r.status_code == 202
    cola.esperar()
    v = c.get(f"/api/trabajos/{id_}/produccion").json()["videos"]["v01"]
    assert v["estado"] == "listo" and v["desactualizado"] is False
    assert not [x for x in v["informe"]["chequeos"] if x["ok"] is False]

    r = c.get(v["archivos"]["mp4"] + "?descargar=1")
    assert r.status_code == 200 and r.headers["content-type"] == "video/mp4" and "attachment" in r.headers["content-disposition"]
    assert c.get(v["archivos"]["vtt"]).text.startswith("WEBVTT")
    # Solo los archivos de la salida, nada más de la carpeta del curso.
    assert c.get(f"/api/trabajos/{id_}/salida/v01/estado.json").status_code == 404
    assert c.get(f"/api/trabajos/{id_}/salida/..%2F..%2Ftrabajo.json/x.mp4").status_code == 404
    assert c.post(f"/api/trabajos/{id_}/producir/no-existe").status_code == 404

    # Si cambia la voz, el video queda marcado como desactualizado.
    c.patch(f"/api/trabajos/{id_}", json={"marca": "riskmann", "voz": "kokoro-alex", "formatos": ["16:9"]})
    assert c.get(f"/api/trabajos/{id_}/produccion").json()["videos"]["v01"]["desactualizado"] is True


def test_producir_con_voz_no_disponible_avisa(datos_copia, monkeypatch):
    from fastapi.testclient import TestClient

    from app.main import app
    monkeypatch.delenv("ELEVENLABS_API_KEY", raising=False)
    c = TestClient(app)
    id_ = _crear(c)
    c.patch(f"/api/trabajos/{id_}", json={"marca": "riskmann", "voz": "carlos", "formatos": ["16:9"]})  # ElevenLabs, sin clave
    assert "Falta la clave de ElevenLabs" in c.get(f"/api/trabajos/{id_}/produccion").json()["voz_falta"]
    r = c.post(f"/api/trabajos/{id_}/producir/v01")
    assert r.status_code == 400 and "ElevenLabs" in r.json()["detail"]
    carlos = next(v for v in c.get("/api/catalogo").json()["voces"] if v["id"] == "carlos")
    assert "ElevenLabs" in carlos["falta"]


def test_env_carga_claves_sin_pisar_las_del_sistema(tmp_path, monkeypatch):
    from app import datos
    (tmp_path / ".env").write_text("# comentario\nCLAVE_A=desde-archivo\nCLAVE_B='con comillas'\nVACIA=\n", encoding="utf-8")
    monkeypatch.delenv("CLAVE_A", raising=False)
    monkeypatch.setenv("CLAVE_B", "del-sistema")
    monkeypatch.delenv("VACIA", raising=False)
    datos._cargar_env(tmp_path / ".env")
    import os
    assert os.environ["CLAVE_A"] == "desde-archivo"
    assert os.environ["CLAVE_B"] == "del-sistema"
    assert "VACIA" not in os.environ
    monkeypatch.delenv("CLAVE_A")
